"""Ask a financial question in plain English, or run metrics for every report.

    python -m src.ask "What was the revenue of TCS in FY 2025-26?"
    python -m src.ask "What was the total income of RBI?"
    python -m src.ask "What was the revenue of TCS from FY 2024-25 to FY 2025-26?"
    python -m src.ask --all                                  # revenue for every report
    python -m src.ask --all revenue,total_income,profit_after_tax
    python -m src.ask --all all                              # every metric
    add -v to see the step-by-step log

--all saves data/processed/extraction_results.csv (and .xlsx when `pip install openpyxl` is done).
The same pipeline answers every question; only the metric (from the question) changes."""
import csv
import logging
import sys

from src.extract.engine import extract_financial_metric
from src.extract.metrics import METRICS
from src.extract.query import label_year, list_reports, parse_question
from src.retrieve.bm25_search import ReportNotFound

log = logging.getLogger("extractor")


def answer_for_year(company, labels, metric, year, scope):
    """Try every report of the company. Prefer the report whose own year is `year`;
    otherwise take the comparative (previous-year) column of a later report."""
    order = sorted(labels, key=lambda l: (label_year(l) != year, -(label_year(l) or 0)))
    first_failure = None
    for label in order:
        r = extract_financial_metric(company, label, metric, year=year, scope=scope)
        if r["status"] in ("OK", "VERIFY"):
            return r
        first_failure = first_failure or r
    return first_failure


def ask(question: str):
    reports = list_reports()
    q = parse_question(question, reports)
    if q["metric"] is None:
        return {"error": "I could not tell which metric you mean. Supported: "
                         + ", ".join(m.replace("_", " ") for m in METRICS)}
    if q["company"] is None:
        names = sorted({c for c, _ in reports})
        return {"error": "I could not tell which company you mean. Reports available: " + ", ".join(names)}
    company, metric, scope = q["company"], q["metric"], q["scope"]
    labels = [l for c, l in reports if c == company]

    years = q["years"]
    if q["years_mode"] == "latest":
        years = [None]                               # latest year of the latest report
        labels = [max(labels, key=lambda l: label_year(l) or 0)]
    elif q["years_mode"] == "all":
        found = set()
        for label in labels:
            found.update(extract_financial_metric(company, label, metric, scope=scope)["year_columns"])
        years = sorted(found)
    results = []
    for y in years:
        r = answer_for_year(company, labels, metric, y, scope)
        if r is not None:
            r["requested_year"] = y
            results.append(r)
    return {"question": question, "parsed": q, "results": results}


def fmt_year(y):
    return f"FY {y - 1}-{str(y)[-2:]}" if y else "latest year"


def print_card(r):
    name = r["metric"].replace("_", " ").title()
    print("-" * 64)
    print(f"Company        : {r['company']}")
    print(f"Metric         : {name}")
    if r["status"] in ("NOT_FOUND", "YEAR_NOT_IN_REPORT"):
        print(f"Value          : NOT FOUND - {r['message']}")
        return
    if r["metric_used"]:
        print(f"Answered with  : {r['metric_used'].replace('_', ' ').title()} (no line is named "
              f"{name.lower()} in this report)")
    print(f"Value          : {r['value']} {r['unit'] or ''} {r['currency'] or ''}".rstrip())
    if r["interpretation"]:
        print(f"Meaning        : {r['interpretation']}")
    print(f"Financial Year : {fmt_year(r['year'])} (from report {r['report']})")
    print(f"Statement      : {r['statement']}")
    print(f"Scope          : {r['scope']}")
    print(f"Source         : Annual Report {r['report']}, page {r['page']}")
    print(f"Matched line   : {r['label']}")
    print("Relevant table :")
    print(f"   columns: {' | '.join(str(y) for y in r['year_columns'])}")
    for line in r["evidence"]:
        print("   " + line)
    print(f"Confidence     : {r['confidence']}")
    if r["status"] == "VERIFY":
        print(f"!! {r['message']}")
    print("Why            :")
    for line in r["reasons"]:
        print("   " + line)
    for w in r["warnings"]:
        if w not in r["message"]:
            print(f"   ! {w}")


# ------------------------------------------------------------------ batch output (CSV / Excel)
COLUMNS = ["company", "report", "metric", "answered_with", "value", "unit", "currency", "value_absolute",
           "year", "statement", "scope", "page", "matched_line", "status", "confidence", "warnings",
           "reasons", "evidence"]


def to_row(r):
    return {"company": r["company"], "report": r["report"], "metric": r["metric"],
            "answered_with": r["metric_used"] or "", "value": r["value"] or "", "unit": r["unit"] or "",
            "currency": r["currency"] or "", "value_absolute": r["value_absolute"],
            "year": r["year"], "statement": r["statement"] or "", "scope": r["scope"] or "",
            "page": r["page"], "matched_line": r["label"] or "", "status": r["status"],
            "confidence": r["confidence"], "warnings": " | ".join(r["warnings"] + ([r["message"]] if r["status"] != "OK" and not r["warnings"] else [])),
            "reasons": "\n".join(r["reasons"]), "evidence": "\n".join(r["evidence"])}


def save_results(rows, base="data/processed/extraction_results"):
    with open(base + ".csv", "w", newline="", encoding="utf-8-sig") as f:      # utf-8-sig: opens cleanly in Excel
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    saved = [base + ".csv"]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("(Excel file skipped: run  python -m pip install openpyxl  to get it)")
        return saved
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row[c] for c in COLUMNS])
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    widths = {"statement": 38, "matched_line": 30, "warnings": 50, "reasons": 60, "evidence": 50}
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(name, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws["H"][1:]:
        cell.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    summary = wb.create_sheet("Summary")
    summary.append(["Measure", "Count"])
    summary.append(["Rows", len(rows)])
    for key in ("status", "confidence"):
        for value in sorted({r[key] for r in rows}):
            summary.append([f"{key} = {value}", sum(r[key] == value for r in rows)])
    summary.column_dimensions["A"].width = 30
    wb.save(base + ".xlsx")
    saved.append(base + ".xlsx")
    return saved


def run_all(metric_arg: str = "revenue"):
    """Chosen metrics for every ingested report: one summary line each, then CSV + Excel."""
    metrics = list(METRICS) if metric_arg == "all" else [m.strip() for m in metric_arg.split(",")]
    unknown = [m for m in metrics if m not in METRICS]
    if unknown:
        sys.exit(f"Unknown metric(s): {', '.join(unknown)}. Available: {', '.join(METRICS)}")
    rows = []
    print(f"{'company':<15}{'report':<8}{'metric':<22}{'value':<16}{'unit':<9}{'year':<6}{'page':<6}{'status':<20}conf")
    for company, report in list_reports():
        for metric in metrics:
            try:
                r = extract_financial_metric(company, report, metric)
            except Exception as e:                    # one bad report/metric must not stop the batch
                log.exception("failed: %s %s %s", company, report, metric)
                r = {"company": company, "report": report, "metric": metric, "metric_used": None,
                     "status": "ERROR", "value": None, "value_float": None, "value_absolute": None,
                     "unit": None, "currency": None, "year": None, "statement": None, "scope": None,
                     "page": None, "label": None, "evidence": [], "confidence": "NONE", "reasons": [],
                     "warnings": [], "message": str(e)}
            rows.append(to_row(r))
            shown = (r["value"] or "-") + ("*" if r["metric_used"] else "")
            print(f"{company:<15}{report:<8}{metric:<22}{shown:<16}{r['unit'] or '-':<9}{str(r['year'] or '-'):<6}"
                  f"{str(r['page'] or '-'):<6}{r['status']:<20}{r['confidence']}")
    if any(r["answered_with"] for r in rows):
        print("* answered with a fallback metric (see the 'answered_with' column)")
    for path in save_results(rows):
        print(f"Saved -> {path}")


def main():
    args = [a for a in sys.argv[1:] if a not in ("-v", "--verbose")]
    logging.basicConfig(level=logging.DEBUG if len(args) != len(sys.argv) - 1 else logging.WARNING,
                        format="%(levelname)s %(message)s")
    if not args:
        sys.exit('Usage: python -m src.ask "What was the revenue of TCS in FY 2025-26?"   or   --all [metrics]')
    try:
        if args[0] == "--all":
            return run_all(args[1] if len(args) > 1 else "revenue")
        out = ask(" ".join(args))
    except ReportNotFound as e:
        sys.exit(str(e))
    except FileNotFoundError:
        sys.exit("No reports have been ingested yet. Run: python -m src.ingest.run_ingest")
    if "error" in out:
        sys.exit(out["error"])
    for r in out["results"]:
        print_card(r)
    if not out["results"]:
        print("No answer found.")


if __name__ == "__main__":
    main()
